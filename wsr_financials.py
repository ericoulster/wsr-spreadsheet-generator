#!/usr/bin/env python3
"""wsr_financials.py - export Wall $treet Raider balance sheets + cash flow statements to Excel.

Reads each entity's financials live over the game's local REST bridge and writes a workbook
(one sheet per entity + a Summary) to a folder. Default scope: YOU (the player) plus every
company you control. Self-contained - it modifies no game files and imports nothing external
beyond `requests` (and `openpyxl` for .xlsx output; it falls back to CSV without it).

    python3 wsr_financials.py --out-dir FOLDER [--symbols RELI,IMD] [--all]
        [--controlled-only] [--player-only] [--overwrite] [--settle 0.3]
        [--runtime PATH | --port N]

The game publishes its bridge port to runtime.json while running; this tool finds it
automatically (Steam/Proton on Linux, %LOCALAPPDATA% on Windows) or via --runtime / --port.
All figures are in $ millions.
"""
import argparse
import csv
import glob
import json
import os
import pathlib
import re
import sys
import time

import requests

APPID = "3525620"


# ----------------------------- bridge client (self-contained) -----------------------------

def _find_runtime():
    """Locate the WSR runtime.json the engine writes while the game is running."""
    home = pathlib.Path.home()
    globs = [
        home / ".steam/*/steamapps/compatdata" / APPID
        / "pfx/drive_c/users/steamuser/AppData/Local/Wall Street Raider/runtime.json",
        home / ".local/share/Steam/steamapps/compatdata" / APPID
        / "pfx/drive_c/users/steamuser/AppData/Local/Wall Street Raider/runtime.json",
    ]
    for g in globs:
        for hit in glob.glob(str(g)):
            return hit
    fixed = []
    la = os.environ.get("LOCALAPPDATA")
    if la:
        fixed.append(pathlib.Path(la) / "Wall Street Raider" / "runtime.json")
    fixed.append(home / "AppData/Local/Wall Street Raider/runtime.json")
    for p in fixed:
        if os.path.exists(p):
            return str(p)
    return None


class Bridge:
    """Minimal REST client for the WSR engine: gamestate + set_view_asset, with port re-read."""

    def __init__(self, runtime=None, port=None, timeout=10.0):
        self.runtime, self.port, self.timeout = runtime, port, timeout
        self._s = requests.Session()
        self._load()

    def _load(self):
        if self.port:
            self.rest = f"http://127.0.0.1:{self.port}"
            return
        rt = self.runtime or _find_runtime()
        if not rt or not os.path.exists(rt):
            raise SystemExit("Could not find WSR runtime.json (is the game running?). "
                             "Pass --runtime PATH or --port N.")
        self.runtime = rt
        self.rest = f"http://127.0.0.1:{json.loads(pathlib.Path(rt).read_text())['rest_port']}"

    def _request(self, method, path, **kw):
        kw.setdefault("timeout", self.timeout)
        try:
            r = self._s.request(method, f"{self.rest}{path}", **kw)
        except requests.ConnectionError:
            self._load()                                   # the engine relaunches on new ports
            r = self._s.request(method, f"{self.rest}{path}", **kw)
        r.raise_for_status()
        return r.json()

    def gamestate(self):
        return self._request("GET", "/gamestate")

    def set_view(self, entity_id):
        return self._request("POST", "/set_view_asset", json={"id": entity_id})


# ----------------------------- statement field maps ($ millions) -----------------------------

COMPANY_ASSETS = [
    ("Cash", "cash"), ("T-Bills", "tBills"), ("Govt Bonds", "govBonds"), ("Corp Bonds", "corpBonds"),
    ("Stock Portfolio", "stocksPortfolioValue"), ("Commodities", "commoditiesPortfolioValue"),
    ("Options", "optPortfolio"), ("Business Loans", "bizLoan"), ("Consumer Loans", "consumerLoan"),
    ("Mortgage Loans", "mortgageLoan"), ("Capital Assets", "capAssets"), ("Goodwill", "goodwill"),
]
COMPANY_LIABS = [
    ("Bank Loan", "loan"), ("Bonds Outstanding", "bondsOut"), ("Demand Deposits", "demandDeposits"),
    ("Certificates of Deposit", "certDeposits"), ("Insurance Reserves", "insurReserves"),
    ("Accrued Tax", "accTax"), ("Cap-Gains Tax Accrued", "capTax"), ("Bad-Debt Reserve", "badDebt"),
    ("Hidden Reserves", "hidReserve"),
]
COMPANY_CF = [
    ("Operating Profit", "operatingProfit"), ("Cash Flow before Debt", "cfBeforeDebt"),
    ("Interest / Debt Service", "@interest"), ("Cash Flow after Debt", "cfAfterDebt"),
    ("Normalized Cash Flow", "normalCashFlo"), ("Est. Cash in 3 Months", "estCashIn3Months"),
    ("Dividend (per share)", "dividend"), ("Commodity Margin P&L", "commodMargin"),
    ("EPS (last 4 qtrs)", "@eps"),
]
PLAYER_ASSETS = [
    ("Cash", "@cash"), ("T-Bills", "@tBills"), ("Stock Portfolio", "stocksPortfolioValue"),
    ("Govt Bonds", "govBondPortfolio"), ("Corp Bonds", "corpBondPortfolio"),
    ("Commodities (MtM)", "commoditiesMtm"), ("Physical Commodities", "physicalCommodValue"),
    ("Options (net)", "optionsNetValue"), ("Advances to Companies", "advancesToCompanies"),
    ("Prepaid Tax", "prepaidTax"),
]
PLAYER_LIABS = [
    ("Income Tax Owed", "incomeTaxOwed"), ("Tax on Corp Shares", "corpSharesTax"),
    ("Wealth Tax (proj.)", "wealthTaxProjected"),
]
PLAYER_CF = [
    ("Annual Net Income", "annualNetIncome"), ("Proj. Annual Cash Flow", "projAnnualCashFlow"),
    ("Living Expenses", "livingExpenses"), ("Non-Cash Expense", "nonCashExpense"),
    ("Realized Cap Gain/Loss", "realizedCapGainLoss"),
]


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return v


def _f(d, k):
    try:
        return float(d.get(k) or 0)
    except (TypeError, ValueError):
        return 0.0


def company_rows(aef, aed, date):
    """rows = [(kind, label, value)]; kind in title/meta/section/sub/line/total/check/blank."""
    rows = [("title", aed.get("name") or "?", aed.get("symbol") or ""),
            ("meta", "As of (game month)", date),
            ("meta", "Market Cap ($M)", num(aed.get("marketCap"))),
            ("meta", "Credit Rating", aed.get("credRating")),
            ("meta", "Mgmt Rating", aed.get("mgmtRating")),
            ("blank", "", ""), ("section", "BALANCE SHEET ($M)", ""), ("sub", "Assets", "")]
    for lbl, k in COMPANY_ASSETS:
        rows.append(("line", lbl, num(aef.get(k))))
    rows += [("total", "Total Assets", num(aef.get("totalAssets"))), ("sub", "Liabilities", "")]
    sum_liab = 0.0
    for lbl, k in COMPANY_LIABS:
        sum_liab += _f(aef, k)
        rows.append(("line", lbl, num(aef.get(k))))
    rows += [("total", "Total Liabilities", num(round(sum_liab, 2))),
             ("line", "  (of which interest-bearing debt)", num(aef.get("totalDebt"))),
             ("sub", "Equity", ""),
             ("total", "Shareholders' Equity", num(aef.get("equity"))),
             ("check", "Balance check: Assets - Liabilities - Equity (~0)",
              num(round(_f(aef, "totalAssets") - sum_liab - _f(aef, "equity"), 2))),
             ("line", "Memo: Line of Credit", num(aef.get("loc"))),
             ("blank", "", ""), ("section", "CASH FLOW STATEMENT ($M)", "")]
    for lbl, k in COMPANY_CF:
        if k == "@interest":
            v = round(_f(aef, "cfBeforeDebt") - _f(aef, "cfAfterDebt"), 2)
        elif k == "@eps":
            v = "   ".join(str(num(aef.get(e))) for e in ("eps1", "eps2", "eps3", "eps4"))
        else:
            v = num(aef.get(k))
        rows.append(("line", lbl, v))
    return rows


def player_rows(apf, top, name, date):
    g = dict(apf or {})
    g["@cash"] = top.get("cash"); g["@tBills"] = top.get("tBills")
    rows = [("title", name or "Player", "PLAYER"),
            ("meta", "As of (game month)", date),
            ("meta", "Net Worth ($M)", num(top.get("netWorth"))),
            ("meta", "Borrow Rate %", num(g.get("borrowRate"))),
            ("blank", "", ""), ("section", "BALANCE SHEET ($M)", ""), ("sub", "Assets", "")]
    for lbl, k in PLAYER_ASSETS:
        rows.append(("line", lbl, num(g.get(k))))
    rows.append(("sub", "Liabilities", ""))
    for lbl, k in PLAYER_LIABS:
        rows.append(("line", lbl, num(g.get(k))))
    rows += [("sub", "Net Worth", ""), ("total", "Net Worth", num(top.get("netWorth"))),
             ("line", "Memo: Line of Credit", num(g.get("lineOfCredit"))),
             ("blank", "", ""), ("section", "CASH FLOW STATEMENT ($M)", "")]
    for lbl, k in PLAYER_CF:
        rows.append(("line", lbl, num(g.get(k))))
    return rows


# ----------------------------- writers -----------------------------

SUMMARY_HEADER = ["Entity", "Symbol", "Market Cap", "Total Assets", "Total Debt",
                  "Equity / Net Worth", "Cash Flow"]


def _sheet_name(name, used):
    base = re.sub(r"[\\/?*\[\]:]", "", str(name))[:28] or "Sheet"
    title, i = base, 1
    while title in used:
        i += 1
        title = f"{base[:25]}_{i}"
    used.add(title)
    return title


def write_xlsx(path, entities, summary):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(SUMMARY_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in summary:
        ws.append(row)
    ws.column_dimensions["A"].width = 30
    used = {"Summary"}
    for ent in entities:
        ws = wb.create_sheet(_sheet_name(ent["sheet"], used))
        for kind, lbl, val in ent["rows"]:
            if kind == "blank":
                ws.append([])
                continue
            ws.append([lbl, val])
            cell = ws.cell(row=ws.max_row, column=1)
            if kind == "title":
                cell.font = Font(bold=True, size=13)
            elif kind == "section":
                cell.font = Font(bold=True, size=11)
            elif kind in ("sub", "total", "check"):
                cell.font = Font(bold=True)
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 18
    wb.save(path)


def write_csv(folder, entities, summary):
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, "summary.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(SUMMARY_HEADER)
        w.writerows(summary)
    used = set()
    for ent in entities:
        fn = re.sub(r"[\\/?*\[\]:]", "", str(ent["sheet"]))[:40] or "entity"
        while fn in used:
            fn += "_"
        used.add(fn)
        with open(os.path.join(folder, f"{fn}.csv"), "w", newline="") as f:
            w = csv.writer(f)
            for kind, lbl, val in ent["rows"]:
                w.writerow([] if kind == "blank" else [lbl, val])


# ----------------------------- main -----------------------------

def main():
    ap = argparse.ArgumentParser(description="Export WSR balance sheets + cash flow statements to Excel.")
    ap.add_argument("--out-dir", required=True, help="folder to write into (created if missing)")
    ap.add_argument("--symbols", default=None, help="comma-separated tickers (overrides scope to just these)")
    ap.add_argument("--all", action="store_true", help="you + every company in the game (slow)")
    ap.add_argument("--controlled-only", action="store_true", help="only the companies you control")
    ap.add_argument("--player-only", action="store_true", help="only your personal statement")
    ap.add_argument("--overwrite", action="store_true", help="single rolling file instead of dated-per-month")
    ap.add_argument("--settle", type=float, default=0.3, help="per-entity read delay, seconds")
    ap.add_argument("--runtime", default=None, help="path to the game's runtime.json (auto-found by default)")
    ap.add_argument("--port", type=int, default=None, help="bridge REST port (bypass runtime.json)")
    a = ap.parse_args()

    b = Bridge(runtime=a.runtime, port=a.port)
    g = b.gamestate()
    yr, mo = g.get("currentYear"), g.get("currentMonth")
    date = f"{yr}-{int(mo):02d}" if yr and mo else "unknown"
    by_sym = {co.get("symbol"): co for co in (g.get("allCompanies") or [])}

    include_player, company_ids = True, []
    if a.symbols:
        include_player = False
        for s in [x.strip() for x in a.symbols.split(",") if x.strip()]:
            co = by_sym.get(s) or by_sym.get(s.upper())
            if co:
                company_ids.append(co["id"])
            else:
                print(f"  warning: symbol '{s}' not found", file=sys.stderr)
    elif a.all:
        company_ids = [co["id"] for co in (g.get("allCompanies") or [])]
    elif a.player_only:
        pass
    elif a.controlled_only:
        include_player = False
        company_ids = [x["id"] for x in (g.get("controlledCompanies") or [])]
    else:
        company_ids = [x["id"] for x in (g.get("controlledCompanies") or [])]
    print(f"WSR financials @ game month {date}: player={include_player}, {len(company_ids)} companies", flush=True)

    entities, summary = [], []
    if include_player:
        b.set_view(2); time.sleep(a.settle); gp = b.gamestate()
        apf = gp.get("activeEntityPlayerFinancials") or {}
        name = gp.get("playerName") or (gp.get("activeEntityData") or {}).get("name") or "Player"
        entities.append({"sheet": "PLAYER", "rows": player_rows(apf, gp, name, date)})
        summary.append([name, "PLAYER", "", "", "", num(gp.get("netWorth")), num(apf.get("projAnnualCashFlow"))])

    for i, cid in enumerate(company_ids):
        b.set_view(cid); time.sleep(a.settle); gc = b.gamestate()
        aef = gc.get("activeEntityFinancials") or {}
        aed = gc.get("activeEntityData") or {}
        sym = aed.get("symbol") or str(cid)
        entities.append({"sheet": sym, "rows": company_rows(aef, aed, date)})
        summary.append([aed.get("name"), sym, num(aed.get("marketCap")), num(aef.get("totalAssets")),
                        num(aef.get("totalDebt")), num(aef.get("equity")), num(aef.get("cfAfterDebt"))])
        if a.all and (i + 1) % 25 == 0:
            print(f"  ...{i + 1}/{len(company_ids)}", flush=True)

    b.set_view(2)  # restore the view to the player

    os.makedirs(a.out_dir, exist_ok=True)
    stamp = "" if a.overwrite else f"_{date}"
    try:
        import openpyxl  # noqa: F401
        path = os.path.join(a.out_dir, f"wsr_financials{stamp}.xlsx")
        write_xlsx(path, entities, summary)
        print(f"wrote {path}  ({len(entities)} entity sheets + Summary)")
    except ImportError:
        folder = os.path.join(a.out_dir, f"wsr_financials{stamp}")
        write_csv(folder, entities, summary)
        print(f"openpyxl not installed -> wrote CSVs to {folder}/  "
              f"({len(entities)} files + summary.csv). `pip install openpyxl` for a single .xlsx workbook.")


if __name__ == "__main__":
    main()
